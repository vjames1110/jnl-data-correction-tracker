"""
The company-wide HR uploads: staff register and muster for many sites
in one file, routed by site code (the HR Department works on every
site).
"""

import io
from datetime import date, timedelta
from decimal import Decimal

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APIClient

from apps.authentication.tests.factories import (
    AdminUserFactory,
    DirectorUserFactory,
    HrDepartmentUserFactory,
    MachineryDepartmentUserFactory,
    ProjectHoUserFactory,
    ProjectManagerUserFactory,
)
from apps.project_monitor.models import (
    StaffDayOverride,
    StaffMember,
)
from apps.project_monitor.services import hr, hr_bulk_import

FORBIDDEN = status.HTTP_403_FORBIDDEN
BAD = status.HTTP_400_BAD_REQUEST


def url(name, *args):
    return reverse(f"project-monitor-api:{name}", args=args)


def today():
    return timezone.localdate()


def days_ago(count):
    return today() - timedelta(days=count)


def previous_month_start():
    return (today().replace(day=1) - timedelta(days=1)).replace(day=1)


@pytest.fixture
def api():
    return APIClient()


@pytest.fixture
def hr_user():
    return HrDepartmentUserFactory()


def workbook_file(rows, name="hr.xlsx"):
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


def upload(api, name, rows, **extra):
    return api.post(
        url(name),
        {"file": workbook_file(rows), **extra},
        format="multipart",
    )


STAFF_HEADER = [
    "Site code",
    "Staff code",
    "Name",
    "Designation",
    "Monthly salary",
    "From date",
    "To date",
    "Effective from",
]
FLAT_HEADER = ["Site code", "Staff code", "Staff name", "Date", "Status"]


def make_staff(site, **overrides):
    values = {
        "name": "Ravi Kumar",
        "staff_code": "E1",
        "monthly_salary": Decimal("30000"),
        "from_date": days_ago(120),
    }
    values.update(overrides)
    return StaffMember.objects.create(site=site, **values)


def staff_upload(api, rows, **extra):
    return upload(
        api, "hr-staff-upload", [STAFF_HEADER, *rows], **extra
    ).data["data"]


# ---- staff register -------------------------------------------------------


@pytest.mark.django_db
def test_one_file_builds_the_register_of_several_sites(
    api, site, other_site, hr_user
):
    api.force_authenticate(user=hr_user)

    result = staff_upload(
        api,
        [
            [site.site_code, "E1", "Ravi Kumar", "Engineer", 42000, days_ago(90), None, None],
            [site.site_code, "E2", "Sita Devi", "Accountant", 30000, days_ago(90), None, None],
            [other_site.site_code, "E1", "Amit Roy", "Store Keeper", 24000, days_ago(30), None, None],
        ],
    )

    assert result["staff_created"] == 3
    assert StaffMember.objects.filter(site=site).count() == 2
    assert StaffMember.objects.filter(site=other_site).count() == 1
    # The same code on two sites is two different people.
    assert StaffMember.objects.get(site=other_site).staff_code == "E1"
    assert {row["site_code"]: row["staff_created"] for row in result["by_site"]} == {
        "CHK": 2,
        "OTH": 1,
    }


@pytest.mark.django_db
def test_reuploading_the_register_changes_nothing(api, site, hr_user):
    api.force_authenticate(user=hr_user)
    rows = [
        [site.site_code, "E1", "Ravi Kumar", "Engineer", 42000, days_ago(90), None, None],
    ]
    staff_upload(api, rows)

    again = staff_upload(api, rows)

    assert again["staff_created"] == 0
    assert again["staff_unchanged"] == 1
    assert StaffMember.objects.count() == 1


@pytest.mark.django_db
def test_a_different_salary_is_refused_without_an_effective_date(
    api, site, hr_user
):
    make_staff(site, monthly_salary=Decimal("30000"))
    api.force_authenticate(user=hr_user)

    result = staff_upload(
        api,
        [[site.site_code, "E1", "Ravi Kumar", "", 35000, days_ago(120), None, None]],
    )

    assert result["salary_conflict"] == 1
    assert "Effective from" in result["errors"][0]
    member = StaffMember.objects.get()
    assert member.monthly_salary == Decimal("30000")
    assert member.to_date is None


@pytest.mark.django_db
def test_a_dated_salary_change_ends_the_old_row_and_starts_a_new_one(
    api, site, hr_user
):
    make_staff(site, monthly_salary=Decimal("30000"), from_date=days_ago(120))
    api.force_authenticate(user=hr_user)
    effective = days_ago(10)

    result = staff_upload(
        api,
        [[site.site_code, "E1", "Ravi Kumar", "", 35000, days_ago(120), None, effective]],
    )

    assert result["staff_revised"] == 1
    old, new = StaffMember.objects.order_by("from_date")
    assert old.monthly_salary == Decimal("30000")
    assert old.to_date == effective - timedelta(days=1)
    assert new.monthly_salary == Decimal("35000")
    assert new.from_date == effective
    assert new.to_date is None
    assert new.staff_code == "E1"
    # Days before the change still cost what they did.
    assert hr.staff_cost_on(old, effective - timedelta(days=2)) == hr.daily_rate(
        old, effective - timedelta(days=2)
    )

    # Uploading the revision file again is a no-op.
    again = staff_upload(
        api,
        [[site.site_code, "E1", "Ravi Kumar", "", 35000, days_ago(120), None, effective]],
    )
    assert again["staff_revised"] == 0
    assert again["staff_unchanged"] == 1
    assert StaffMember.objects.count() == 2


@pytest.mark.django_db
def test_designation_leaving_date_and_a_first_code_are_filled_in(
    api, site, hr_user
):
    make_staff(site, staff_code="", designation="")
    api.force_authenticate(user=hr_user)
    left = days_ago(5)

    result = staff_upload(
        api,
        [[site.site_code, "E9", "Ravi Kumar", "Engineer", 30000, days_ago(120), left, None]],
    )

    assert result["staff_updated"] == 1
    member = StaffMember.objects.get()
    assert (member.staff_code, member.designation, member.to_date) == (
        "E9",
        "Engineer",
        left,
    )


@pytest.mark.django_db
def test_bad_register_rows_are_reported_without_stopping(
    api, site, hr_user
):
    api.force_authenticate(user=hr_user)

    result = staff_upload(
        api,
        [
            ["NOPE", "E1", "Ghost", "", 1000, days_ago(9), None, None],
            [site.site_code, "E2", "", "", 1000, days_ago(9), None, None],
            [site.site_code, "E3", "No Date", "", 1000, "soon", None, None],
            [site.site_code, "E4", "Negative", "", -5, days_ago(9), None, None],
            [site.site_code, "E5", "Back To Front", "", 1000, days_ago(2), days_ago(9), None],
            [site.site_code, "E6", "Good One", "", 1000, days_ago(9), None, None],
        ],
    )

    assert result["site_not_found"] == 1
    assert result["invalid"] == 4
    assert result["staff_created"] == 1
    assert result["error_count"] == 5
    assert StaffMember.objects.get().name == "Good One"


@pytest.mark.django_db
def test_the_register_file_needs_a_recognisable_header(api, site, hr_user):
    api.force_authenticate(user=hr_user)

    response = upload(
        api, "hr-staff-upload", [["Foo", "Bar"], ["1", "2"]]
    )

    assert response.status_code == BAD


@pytest.mark.django_db
@pytest.mark.parametrize(
    "uploader_factory",
    [
        ProjectManagerUserFactory,
        MachineryDepartmentUserFactory,
        ProjectHoUserFactory,
    ],
)
def test_only_hr_can_upload_a_register(api, site, uploader_factory):
    api.force_authenticate(user=uploader_factory())

    response = upload(
        api,
        "hr-staff-upload",
        [
            STAFF_HEADER,
            [site.site_code, "E1", "Ravi", "", 1000, days_ago(9), None, None],
        ],
    )

    if response.status_code == 200:
        assert response.data["data"]["not_permitted"] == 1
    else:
        assert response.status_code == FORBIDDEN
    assert StaffMember.objects.count() == 0


@pytest.mark.django_db
def test_admin_can_upload_a_register_too(api, site, other_site):
    api.force_authenticate(user=AdminUserFactory())

    result = staff_upload(
        api,
        [
            [site.site_code, "E1", "A", "", 1000, days_ago(9), None, None],
            [other_site.site_code, "E1", "B", "", 1000, days_ago(9), None, None],
        ],
    )

    assert result["staff_created"] == 2


# ---- muster: flat rows ------------------------------------------------------


def muster_upload(api, rows, header=FLAT_HEADER, **extra):
    return upload(
        api, "hr-muster-upload", [header, *rows], **extra
    ).data["data"]


@pytest.mark.django_db
def test_a_flat_muster_lands_on_each_site(api, site, other_site, hr_user):
    ravi = make_staff(site)
    amit = make_staff(other_site, name="Amit Roy", staff_code="A1")
    api.force_authenticate(user=hr_user)
    absent_day, half_day = days_ago(5), days_ago(4)

    result = muster_upload(
        api,
        [
            [site.site_code, "E1", "Ravi Kumar", absent_day, "A"],
            [site.site_code, "E1", "Ravi Kumar", half_day, "HD"],
            [site.site_code, "E1", "Ravi Kumar", days_ago(3), "P"],
            [other_site.site_code, "A1", "Amit Roy", absent_day, "Absent"],
        ],
    )

    assert result["muster_absent"] == 2
    assert result["muster_half"] == 1
    assert result["muster_unchanged"] == 1
    assert StaffDayOverride.objects.get(staff=ravi, date=absent_day).amount == 0
    assert StaffDayOverride.objects.get(staff=ravi, date=half_day).amount == (
        hr.daily_rate(ravi, half_day) / 2
    ).quantize(Decimal("0.01"))
    assert StaffDayOverride.objects.get(staff=amit, date=absent_day).amount == 0
    assert not StaffDayOverride.objects.filter(staff=ravi, date=days_ago(3)).exists()
    assert {row["site_code"] for row in result["by_site"]} == {"CHK", "OTH"}


@pytest.mark.django_db
def test_reuploading_a_muster_changes_nothing(api, site, hr_user):
    make_staff(site)
    api.force_authenticate(user=hr_user)
    rows = [
        [site.site_code, "E1", "Ravi Kumar", days_ago(5), "A"],
        [site.site_code, "E1", "Ravi Kumar", days_ago(4), "HD"],
    ]
    muster_upload(api, rows)

    again = muster_upload(api, rows)

    assert again["muster_absent"] == 0
    assert again["muster_half"] == 0
    assert again["muster_unchanged"] == 2
    assert StaffDayOverride.objects.count() == 2


@pytest.mark.django_db
def test_present_clears_a_muster_day_but_never_a_manual_one(
    api, site, hr_user
):
    ravi = make_staff(site)
    api.force_authenticate(user=hr_user)
    corrected, manual = days_ago(5), days_ago(4)
    muster_upload(api, [[site.site_code, "E1", "Ravi", corrected, "A"]])
    hr.set_override(
        staff=ravi, day=manual, amount=Decimal("777"), note="Overtime"
    )

    result = muster_upload(
        api,
        [
            [site.site_code, "E1", "Ravi", corrected, "P"],
            [site.site_code, "E1", "Ravi", manual, "P"],
        ],
    )

    assert result["muster_cleared"] == 1
    assert result["muster_kept_manual"] == 1
    assert not StaffDayOverride.objects.filter(staff=ravi, date=corrected).exists()
    assert StaffDayOverride.objects.get(staff=ravi, date=manual).amount == 777


@pytest.mark.django_db
def test_muster_problems_are_reported_row_by_row(api, site, hr_user):
    make_staff(site, from_date=days_ago(10))
    api.force_authenticate(user=hr_user)

    result = muster_upload(
        api,
        [
            [site.site_code, "E1", "Ravi", days_ago(2), "A"],
            [site.site_code, "ZZ", "Nobody", days_ago(2), "A"],
            [site.site_code, "E1", "Ravi", days_ago(2), "MAYBE"],
            [site.site_code, "E1", "Ravi", today() + timedelta(days=2), "A"],
            [site.site_code, "E1", "Ravi", days_ago(50), "A"],
            ["NOPE", "E1", "Ravi", days_ago(2), "A"],
        ],
    )

    assert result["muster_absent"] == 1
    assert result["staff_not_found"] == 1
    assert result["off_payroll"] == 1
    assert result["site_not_found"] == 1
    assert result["invalid"] == 2
    assert StaffDayOverride.objects.count() == 1


@pytest.mark.django_db
def test_muster_matches_a_name_when_the_register_has_no_code(
    api, site, hr_user
):
    make_staff(site, staff_code="")
    api.force_authenticate(user=hr_user)

    result = muster_upload(
        api, [[site.site_code, "", "ravi kumar", days_ago(3), "A"]]
    )

    assert result["muster_absent"] == 1


@pytest.mark.django_db
@pytest.mark.parametrize(
    "uploader_factory",
    [
        ProjectHoUserFactory,
        ProjectManagerUserFactory,
        MachineryDepartmentUserFactory,
    ],
)
def test_only_hr_can_upload_a_muster(api, site, uploader_factory):
    make_staff(site)
    api.force_authenticate(user=uploader_factory())

    response = upload(
        api,
        "hr-muster-upload",
        [FLAT_HEADER, [site.site_code, "E1", "Ravi", days_ago(3), "A"]],
    )

    if response.status_code == 200:
        assert response.data["data"]["not_permitted"] == 1
    else:
        assert response.status_code == FORBIDDEN
    assert StaffDayOverride.objects.count() == 0


# ---- muster: month grid --------------------------------------------------------


def grid_header(days, first=("Site code", "Staff code", "Staff name")):
    return [*first, *days]


@pytest.mark.django_db
def test_a_grid_with_real_date_headers(api, site, other_site, hr_user):
    ravi = make_staff(site)
    amit = make_staff(other_site, name="Amit Roy", staff_code="A1")
    api.force_authenticate(user=hr_user)
    d1, d2, d3 = days_ago(6), days_ago(5), days_ago(4)

    result = upload(
        api,
        "hr-muster-upload",
        [
            grid_header([d1, d2, d3]),
            [site.site_code, "E1", "Ravi", "P", "A", "HD"],
            [other_site.site_code, "A1", "Amit", "P", "P", "A"],
        ],
    ).data["data"]

    assert result["muster_absent"] == 2
    assert result["muster_half"] == 1
    assert StaffDayOverride.objects.get(staff=ravi, date=d2).amount == 0
    assert StaffDayOverride.objects.get(staff=amit, date=d3).amount == 0


@pytest.mark.django_db
def test_a_grid_of_day_numbers_uses_the_month_named_in_the_sheet(
    api, site, hr_user
):
    start = previous_month_start()
    ravi = make_staff(site, from_date=start - timedelta(days=30))
    api.force_authenticate(user=hr_user)

    result = upload(
        api,
        "hr-muster-upload",
        [
            [f"Muster for {start:%B %Y}"],
            grid_header([1, 2, 3, 4]),
            [site.site_code, "E1", "Ravi", "P", "A", None, "HD"],
            [None, None, "Total"],
        ],
    ).data["data"]

    assert result["muster_absent"] == 1
    assert result["muster_half"] == 1
    assert StaffDayOverride.objects.get(
        staff=ravi, date=start.replace(day=2)
    ).amount == 0
    assert StaffDayOverride.objects.filter(staff=ravi).count() == 2


@pytest.mark.django_db
def test_a_grid_of_day_numbers_can_take_the_month_from_the_upload(
    api, site, hr_user
):
    start = previous_month_start()
    make_staff(site, from_date=start - timedelta(days=30))
    api.force_authenticate(user=hr_user)

    result = upload(
        api,
        "hr-muster-upload",
        [grid_header([1, 2]), [site.site_code, "E1", "Ravi", "A", "A"]],
        month=f"{start:%Y-%m}",
    ).data["data"]

    assert result["muster_absent"] == 2


@pytest.mark.django_db
def test_a_grid_of_day_numbers_with_no_month_is_a_clear_error(
    api, site, hr_user
):
    make_staff(site)
    api.force_authenticate(user=hr_user)

    response = upload(
        api,
        "hr-muster-upload",
        [grid_header([1, 2]), [site.site_code, "E1", "Ravi", "A", "A"]],
    )

    assert response.status_code == BAD
    assert "month" in str(response.data).lower()


@pytest.mark.django_db
def test_the_current_months_future_days_are_left_blank_not_errors(
    api, site, hr_user
):
    make_staff(site, from_date=today().replace(day=1) - timedelta(days=30))
    api.force_authenticate(user=hr_user)
    month_start = today().replace(day=1)

    result = upload(
        api,
        "hr-muster-upload",
        [
            grid_header([month_start + timedelta(days=i) for i in range(0, 3)]),
            [site.site_code, "E1", "Ravi", None, None, None],
        ],
    ).data["data"]

    assert result["error_count"] == 0


@pytest.mark.django_db
def test_an_unreadable_muster_layout_is_a_clear_error(api, site, hr_user):
    api.force_authenticate(user=hr_user)

    response = upload(
        api, "hr-muster-upload", [["Foo", "Bar"], ["1", "2"]]
    )

    assert response.status_code == BAD


# ---- status words, months, templates ----------------------------------------------


@pytest.mark.parametrize(
    "value, expected",
    [
        ("P", "P"), ("present", "P"), ("WO", "P"), ("PL", "P"), (1, "P"), (1.0, "P"),
        ("HD", "H"), ("Half Day", "H"), (0.5, "H"),
        ("A", "A"), ("ab", "A"), ("LWP", "A"), (0, "A"),
        ("", None), (None, None), ("   ", None),
        ("??", "?"), ("maybe", "?"),
    ],
)
def test_status_words(value, expected):
    assert hr_bulk_import.parse_status(value) == expected


@pytest.mark.parametrize(
    "text, expected",
    [
        ("September 2026", (2026, 9)),
        ("Sep-26", (2026, 9)),
        ("Muster for Sept 2026", (2026, 9)),
        ("2026-09", (2026, 9)),
        ("09/2026", (2026, 9)),
        ("Attendance", None),
        ("Site code", None),
    ],
)
def test_month_in_text(text, expected):
    assert hr_bulk_import._month_in_text(text) == expected


@pytest.mark.django_db
def test_templates_download_and_read_back(api, hr_user, site):
    api.force_authenticate(user=hr_user)

    for name, query in (
        ("hr-staff-template", {}),
        ("hr-muster-template", {"layout": "flat"}),
        ("hr-muster-template", {"layout": "grid", "month": "2026-02"}),
    ):
        response = api.get(url(name), query)
        assert response.status_code == 200
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        assert "Site codes" in book.sheetnames
        assert site.site_code in [
            row[0] for row in book["Site codes"].iter_rows(values_only=True)
        ]

    grid = openpyxl.load_workbook(
        io.BytesIO(
            api.get(
                url("hr-muster-template"),
                {"layout": "grid", "month": "2026-02"},
            ).content
        )
    )["Muster"]
    header = [c for c in list(grid.iter_rows(values_only=True))[1]]
    assert header[3:] == list(range(1, 29))


@pytest.mark.django_db
def test_the_filled_in_flat_template_round_trips(api, hr_user, site):
    make_staff(site, staff_code="EMP001", name="Site Engineer name", from_date=date(2020, 1, 1))
    api.force_authenticate(user=hr_user)
    content = api.get(
        url("hr-muster-template"),
        {"layout": "flat", "month": f"{previous_month_start():%Y-%m}", "site": str(site.id)},
    ).content

    response = api.post(
        url("hr-muster-upload"),
        {"file": SimpleUploadedFile("t.xlsx", content)},
        format="multipart",
    )

    assert response.status_code == 200, response.data
    assert response.data["data"]["muster_absent"] == 1
    assert response.data["data"]["muster_half"] == 1


@pytest.mark.django_db
def test_staff_code_is_shown_and_editable_on_the_register(api, site, hr_user):
    api.force_authenticate(user=hr_user)

    created = api.post(
        url("hr-staff-list"),
        {
            "site": str(site.id),
            "staff_code": "E77",
            "name": "New Person",
            "monthly_salary": "20000",
            "from_date": days_ago(5).isoformat(),
        },
        format="json",
    )

    assert created.status_code == 200, created.data
    assert created.data["data"]["staff_code"] == "E77"
    duplicate = api.post(
        url("hr-staff-list"),
        {
            "site": str(site.id),
            "staff_code": "E77",
            "name": "Someone Else",
            "monthly_salary": "20000",
            "from_date": days_ago(5).isoformat(),
        },
        format="json",
    )
    assert duplicate.status_code == BAD
