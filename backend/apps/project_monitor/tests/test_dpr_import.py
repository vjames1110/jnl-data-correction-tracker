import io
from datetime import timedelta
from decimal import Decimal

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from rest_framework.exceptions import ValidationError

from apps.project_monitor.models import (
    DprEntry,
    DprEntrySource,
    DprItem,
)
from apps.project_monitor.services import dpr_import
from apps.project_monitor.tests.finance_helpers import (
    days_ago,
    make_item,
)


def xlsx_upload(rows, name="upload.xlsx"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


def csv_upload(text, name="upload.csv"):
    return SimpleUploadedFile(
        name, text.encode("utf-8"), content_type="text/csv"
    )


ITEM_ROWS = [
    ["Schedule of quantities"],
    [],
    ["Item no", "Description", "Unit", "Qty", "Rate"],
    ["1.1", "Earthwork in embankment", "cum", 1000, 250],
    ["1.2", "PCC M15", "cum", "500", "Rs. 4,500"],
    ["1.3", "No rate item", "nos", 5, 0],
    ["", "", "", "", ""],
]


@pytest.mark.django_db
def test_item_import_finds_the_header_and_creates_valid_rows(
    site, pm
):
    result = dpr_import.import_items(
        site, xlsx_upload(ITEM_ROWS), actor=pm
    )

    assert result["created"] == 2
    assert result["skipped_invalid"] == 1
    assert "Row 6" in result["errors"][0]
    items = {i.item_no: i for i in DprItem.objects.filter(site=site)}
    assert items["1.2"].rate == Decimal("4500")
    assert items["1.2"].scope_qty == Decimal("500")
    assert [i.row_order for i in items.values()] == [1, 2]


@pytest.mark.django_db
def test_item_import_is_idempotent(site, pm):
    dpr_import.import_items(site, xlsx_upload(ITEM_ROWS), actor=pm)

    again = dpr_import.import_items(
        site, xlsx_upload(ITEM_ROWS), actor=pm
    )

    assert again["created"] == 0
    assert again["skipped_existing"] == 2
    assert DprItem.objects.filter(site=site).count() == 2


@pytest.mark.django_db
def test_item_import_matches_by_description_when_no_item_no(
    site, pm
):
    make_item(site, item_no="", description="Earthwork in embankment")

    result = dpr_import.import_items(
        site,
        xlsx_upload(
            [
                ["Description", "Qty", "Rate"],
                ["earthwork in embankment", 10, 10],
            ]
        ),
        actor=pm,
    )

    assert result["created"] == 0
    assert result["skipped_existing"] == 1


@pytest.mark.django_db
def test_item_import_works_from_csv(site, pm):
    result = dpr_import.import_items(
        site,
        csv_upload(
            "Item No,Description,Unit,Quantity,Rate\n"
            "9.1,Steel,MT,10,60000\n"
        ),
        actor=pm,
    )

    assert result["created"] == 1


@pytest.mark.django_db
def test_import_rejects_bad_files(site, pm):
    with pytest.raises(ValidationError) as legacy:
        dpr_import.import_items(
            site, xlsx_upload([["x"]], name="old.xls")
        )
    assert "xlsx" in str(legacy.value.detail).lower()

    with pytest.raises(ValidationError):
        dpr_import.import_items(
            site, csv_upload("a,b\n1,2\n", name="notes.txt")
        )
    with pytest.raises(ValidationError):
        dpr_import.import_items(
            site, xlsx_upload([["Foo", "Bar"], [1, 2]])
        )
    with pytest.raises(ValidationError):
        dpr_import.import_items(
            site,
            SimpleUploadedFile("broken.xlsx", b"not a workbook"),
        )


DPR_HEADER = [
    "Date (YYYY-MM-DD)",
    "Item no",
    "Description",
    "Unit",
    "Qty done",
    "Location / chainage",
    "Agency",
    "Remarks",
]


@pytest.mark.django_db
def test_dpr_upload_creates_entries_with_rate_snapshots(site, pm):
    item = make_item(site)
    today = timezone.localdate()

    result = dpr_import.import_dpr_entries(
        site,
        xlsx_upload(
            [
                DPR_HEADER,
                [today, "1.1", "", "cum", 40, "Ch 12", "ABC", "ok"],
                [
                    (today - timedelta(days=1)).strftime(
                        "%d-%m-%Y"
                    ),
                    "",
                    "earthwork in embankment",
                    "cum",
                    10,
                    "",
                    "",
                    "",
                ],
            ]
        ),
        actor=pm,
    )

    assert result["created"] == 2
    entries = DprEntry.objects.filter(item=item)
    assert {e.source for e in entries} == {DprEntrySource.EXCEL}
    assert {e.rate for e in entries} == {Decimal("100")}
    assert entries.get(date=today).location == "Ch 12"


@pytest.mark.django_db
def test_dpr_upload_is_idempotent_but_keeps_genuine_repeats(
    site, pm
):
    make_item(site)
    today = timezone.localdate()
    rows = [
        DPR_HEADER,
        [today, "1.1", "", "", 40, "Ch 12", "", ""],
        [today, "1.1", "", "", 40, "Ch 12", "", ""],
    ]

    first = dpr_import.import_dpr_entries(
        site, xlsx_upload(rows), actor=pm
    )
    second = dpr_import.import_dpr_entries(
        site, xlsx_upload(rows), actor=pm
    )

    assert first["created"] == 2
    assert second["created"] == 0
    assert second["skipped_duplicate"] == 2
    assert DprEntry.objects.count() == 2


@pytest.mark.django_db
def test_dpr_upload_reports_locked_unknown_and_invalid_rows(site, pm):
    make_item(site)
    today = timezone.localdate()

    result = dpr_import.import_dpr_entries(
        site,
        xlsx_upload(
            [
                DPR_HEADER,
                [days_ago(15), "1.1", "", "", 5, "", "", ""],
                [today, "9.9", "Nope", "", 5, "", "", ""],
                ["not a date", "1.1", "", "", 5, "", "", ""],
                [today, "1.1", "", "", 0, "", "", ""],
                [today, "1.1", "", "", 7, "", "", ""],
            ]
        ),
        actor=pm,
    )

    assert result["created"] == 1
    assert result["skipped_locked"] == 1
    assert result["item_not_found"] == 1
    assert result["invalid"] == 2


@pytest.mark.django_db
def test_dpr_upload_needs_date_qty_and_an_item_column(site, pm):
    with pytest.raises(ValidationError):
        dpr_import.import_dpr_entries(
            site,
            xlsx_upload([["Date", "Qty"], ["2026-01-01", 5]]),
        )


@pytest.mark.django_db
def test_templates_are_valid_workbooks(site):
    make_item(site)

    dpr_bytes = dpr_import.build_dpr_template(site)
    items_bytes = dpr_import.build_item_template()

    dpr_sheet = openpyxl.load_workbook(io.BytesIO(dpr_bytes)).active
    assert dpr_sheet.max_row == 2
    assert dpr_sheet["B2"].value == "1.1"
    items_sheet = openpyxl.load_workbook(io.BytesIO(items_bytes)).active
    assert items_sheet["B1"].value == "Description"
