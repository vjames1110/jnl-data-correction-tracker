"""
Excel/CSV import for the DPR: the contract item list, and daily
quantities. Server-side via ``openpyxl`` (``.xlsx``) or the stdlib
``csv`` module - legacy ``.xls`` is rejected with a clear message.

Both imports are idempotent: re-uploading a file does not duplicate
items or DPR entries (the prototype re-appended every row).
"""

import csv
import io
import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from django.core.exceptions import (
    ValidationError as DjangoValidationError,
)
from django.db import transaction
from django.utils import timezone
from rest_framework.exceptions import ValidationError

from apps.project_monitor.models import (
    DprEntry,
    DprEntrySource,
    DprItem,
)
from apps.project_monitor.services import boq
from apps.project_monitor.services import dpr as dpr_service

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
HEADER_SCAN_ROWS = 30
MAX_REPORTED_ERRORS = 20
ZERO = Decimal("0")

# A header that names an amount / value is never a rate column.
_NOT_AMOUNT = r"(?!.*(amount|value|total))"
ITEM_PATTERNS = [
    # Most specific first: a header goes to the first key it matches.
    (
        "authority",
        rf"^{_NOT_AMOUNT}(?=.*rate)"
        r".*(authority|estimat|schedule|sor\b|railway)",
    ),
    ("percent", r"%|percent|(above|below)\b|premium"),
    ("desc", r"desc|particular|item of work|name of work"),
    ("unit", r"^unit|uom"),
    ("qty", r"qty|quantity"),
    ("rate", rf"^{_NOT_AMOUNT}.*rate"),
    ("no", r"^(item|s\.?\s*no|sl|sr|item\s*no)"),
]
DPR_PATTERNS = [
    ("date", r"date"),
    ("desc", r"desc"),
    ("qty", r"qty|quantity"),
    ("location", r"location|chainage"),
    ("agency", r"agency"),
    ("remark", r"remark"),
    ("no", r"item"),
]
DATE_FORMATS = (
    "%Y-%m-%d",
    "%d-%m-%Y",
    "%d/%m/%Y",
    "%d.%m.%Y",
)


def read_rows(uploaded_file) -> list[list]:
    """First sheet of an ``.xlsx``, or a ``.csv``, as raw rows."""
    name = (getattr(uploaded_file, "name", "") or "").lower()
    if name.endswith(".xls"):
        raise ValidationError(
            {
                "file": (
                    "Old .xls files are not supported - "
                    "save the sheet as .xlsx or .csv."
                )
            }
        )
    if not name.endswith((".xlsx", ".csv")):
        raise ValidationError(
            {"file": "Upload an .xlsx or .csv file."}
        )
    if uploaded_file.size > MAX_UPLOAD_BYTES:
        raise ValidationError(
            {"file": "The file is larger than 5 MB."}
        )

    content = uploaded_file.read()
    try:
        if name.endswith(".csv"):
            text = content.decode("utf-8-sig")
            rows = [list(row) for row in csv.reader(io.StringIO(text))]
        else:
            workbook = openpyxl.load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
            )
            rows = [
                list(row)
                for row in workbook.worksheets[0].iter_rows(
                    values_only=True
                )
            ]
    except (UnicodeDecodeError, csv.Error) as exc:
        raise ValidationError(
            {"file": "The file could not be read."}
        ) from exc
    except Exception as exc:  # openpyxl raises many types
        raise ValidationError(
            {"file": "The file is not a valid spreadsheet."}
        ) from exc

    if len(rows) > MAX_ROWS:
        raise ValidationError(
            {
                "file": (
                    f"The file has more than {MAX_ROWS} rows."
                )
            }
        )
    return rows


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


_NUMBER_RE = re.compile(r"-?\d[\d,]*\.?\d*|-?\.\d+")


def _number(value) -> Decimal:
    """
    A cell as a Decimal: real numbers pass through; text like
    ``"Rs. 4,500"`` yields its first number (4500); anything with no
    number is zero.
    """
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    match = _NUMBER_RE.search(_text(value))
    if not match:
        return ZERO
    try:
        return Decimal(match.group().replace(",", ""))
    except InvalidOperation:
        return ZERO


def _qty(value) -> Decimal:
    return _number(value).quantize(Decimal("0.001"))


def _money(value) -> Decimal:
    return _number(value).quantize(Decimal("0.01"))


def _to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _detect_columns(row, patterns) -> dict:
    columns = {}
    for index, cell in enumerate(row):
        header = _text(cell).lower()
        if not header:
            continue
        for key, pattern in patterns:
            if key not in columns and re.search(
                pattern, header
            ):
                columns[key] = index
                break
    return columns


def _find_header(rows, patterns, required):
    for row_index, row in enumerate(
        rows[:HEADER_SCAN_ROWS]
    ):
        columns = _detect_columns(row, patterns)
        if all(key in columns for key in required):
            return row_index, columns
    return None, None


def _cell(row, columns, key):
    index = columns.get(key)
    if index is None or index >= len(row):
        return None
    return row[index]


def _record_error(errors, row_number, message):
    if len(errors) < MAX_REPORTED_ERRORS:
        errors.append(f"Row {row_number}: {message}")


def _item_key(item_no: str) -> str:
    """An item number as used for the group tree: ``4.1.`` -> ``4.1``."""
    return item_no.strip().rstrip(".")


def _parent_keys(key: str) -> list[str]:
    """``4.1.2`` -> ``["4.1", "4"]``: the numbers of the groups above."""
    parts = key.split(".")
    return [
        ".".join(parts[:end]) for end in range(len(parts) - 1, 0, -1)
    ]


def _percent(value):
    """
    A tender percentage cell: ``5.25``, ``-3``, ``5.25%`` or
    ``Below 3.5%``. Blank stays ``None`` (use the contract-wide %).
    """
    text = _text(value)
    if not text:
        return None
    number = _number(value)
    lowered = text.lower()
    if "below" in lowered or "less" in lowered:
        number = -abs(number)
    elif "above" in lowered or "premium" in lowered:
        number = abs(number)
    return number.quantize(Decimal("0.001"))


def _find_group(by_no, key):
    """
    The nearest group above ``key`` an item may sit under: it must
    exist, be a group and leave room for another level. ``None`` =
    top level (also what a flat list gets, as before).
    """
    for parent_key in _parent_keys(key):
        candidate = by_no.get(parent_key)
        if (
            candidate is not None
            and candidate.is_heading
            and candidate.level() < DprItem.MAX_LEVEL
        ):
            return candidate
    return None


@transaction.atomic
def import_items(site, uploaded_file, actor=None) -> dict:
    """
    Import a Railway BOQ / item list. Columns found by name: Item no,
    Description, Unit, Qty, and any of Authority rate, Tender %
    (above / below) and Quoted / bid rate. Rows nest by item number
    (4 -> 4.1 -> 4.1.1); a numbered row with no quantity or rate that
    has rows nested under it becomes a group. Idempotent: an item
    already on the site (same item no, else same description) is
    skipped.
    """
    rows = read_rows(uploaded_file)
    header_index, columns = _find_header(
        rows, ITEM_PATTERNS, ("desc", "qty")
    )
    if header_index is not None and not (
        "rate" in columns or "authority" in columns
    ):
        header_index = None
    if header_index is None:
        raise ValidationError(
            {
                "file": (
                    "Could not find a header row with "
                    "Description, Qty and Rate columns."
                )
            }
        )

    body = [
        (offset, row)
        for offset, row in enumerate(
            rows[header_index + 1 :], start=header_index + 2
        )
        if any(_text(cell) for cell in row)
    ]

    # Every group number that has something nested under it.
    group_keys = set()
    for _offset, row in body:
        key = _item_key(_text(_cell(row, columns, "no")))
        if key:
            group_keys.update(_parent_keys(key))

    existing = list(DprItem.objects.filter(site=site))
    by_no = {
        _item_key(item.item_no): item
        for item in existing
        if item.item_no
    }
    by_desc = {
        item.description.lower(): item for item in existing
    }

    created = 0
    headings = 0
    skipped_existing = 0
    skipped_invalid = 0
    errors = []
    notes = []
    for offset, row in body:
        description = _text(_cell(row, columns, "desc"))
        item_no = _text(_cell(row, columns, "no"))
        key = _item_key(item_no)
        qty = _qty(_cell(row, columns, "qty"))
        rate = _money(_cell(row, columns, "rate"))
        authority = _money(_cell(row, columns, "authority"))
        percent = _percent(_cell(row, columns, "percent"))

        is_heading = bool(
            description
            and key in group_keys
            and qty <= 0
            and rate <= 0
            and authority <= 0
        )
        if not is_heading and (
            not description
            or qty <= 0
            or (rate <= 0 and authority <= 0)
        ):
            skipped_invalid += 1
            _record_error(
                errors,
                offset,
                "needs a description, a quantity above "
                "zero and a rate above zero.",
            )
            continue

        if (item_no and key in by_no) or (
            not item_no and description.lower() in by_desc
        ):
            skipped_existing += 1
            continue

        fields = {}
        if not is_heading:
            authority_rate = authority if authority > 0 else None
            tender_percent = percent
            if (
                authority_rate is not None
                and percent is None
                and rate > 0
            ):
                # Only the quoted rate was given: work back to the
                # percentage that produces it.
                tender_percent = (
                    (rate / authority_rate - 1) * 100
                ).quantize(Decimal("0.001"))
            if tender_percent is not None and not (
                Decimal("-100") <= tender_percent <= Decimal("500")
            ):
                skipped_invalid += 1
                _record_error(
                    errors,
                    offset,
                    "the tender percentage must be between "
                    "-100 and 500.",
                )
                continue
            if (
                authority_rate is not None
                and percent is None
                and rate > 0
            ):
                bid = (
                    authority_rate * (1 + tender_percent / 100)
                ).quantize(Decimal("0.01"))
                if bid != rate:
                    notes.append(
                        f"Row {offset}: the quoted rate {rate} "
                        f"became {bid} after rounding the "
                        "percentage."
                    )
            fields = {
                "scope_qty": qty,
                "rate": rate,
                "authority_rate": authority_rate,
                "tender_percent": tender_percent,
            }

        try:
            item = dpr_service.create_item(
                site=site,
                description=description,
                item_no=item_no,
                unit=_text(_cell(row, columns, "unit")),
                parent=(
                    _find_group(by_no, key) if key else None
                ),
                is_heading=is_heading,
                actor=actor,
                **fields,
            )
        except DjangoValidationError as exc:
            skipped_invalid += 1
            _record_error(
                errors, offset, "; ".join(exc.messages)
            )
            continue
        if item_no:
            by_no[key] = item
        by_desc[item.description.lower()] = item
        created += 1
        headings += 1 if is_heading else 0

    return {
        "created": created,
        "headings": headings,
        "skipped_existing": skipped_existing,
        "skipped_invalid": skipped_invalid,
        "errors": errors,
        "notes": notes[:MAX_REPORTED_ERRORS],
    }


@transaction.atomic
def import_dpr_entries(
    site, uploaded_file, actor=None, today=None
) -> dict:
    today = today or timezone.localdate()
    rows = read_rows(uploaded_file)
    header_index, columns = _find_header(
        rows, DPR_PATTERNS, ("date", "qty")
    )
    if header_index is None or (
        "no" not in columns and "desc" not in columns
    ):
        raise ValidationError(
            {
                "file": (
                    "Could not find a header row with Date, "
                    "Qty and an Item no or Description "
                    "column."
                )
            }
        )

    items = list(DprItem.objects.filter(site=site))
    by_no = {
        item.item_no: item for item in items if item.item_no
    }
    by_desc = {item.description.lower(): item for item in items}
    series = boq.escalation_series(site)

    parsed = []
    for offset, row in enumerate(
        rows[header_index + 1 :], start=header_index + 2
    ):
        if not any(_text(cell) for cell in row):
            continue
        parsed.append((offset, row))

    # Existing Excel-sourced rows, counted per natural key, so an
    # identical re-upload is skipped one-for-one (two genuinely
    # identical lines in one file are still both kept the first time).
    existing_counter = Counter(
        (
            entry.item_id,
            entry.date,
            entry.qty,
            entry.location,
            entry.agency,
            entry.remarks,
        )
        for entry in DprEntry.objects.filter(
            item__site=site, source=DprEntrySource.EXCEL
        )
    )

    created = 0
    skipped_duplicate = 0
    skipped_locked = 0
    item_not_found = 0
    invalid = 0
    errors = []
    locked_cache = {}

    for offset, row in parsed:
        day = _to_date(_cell(row, columns, "date"))
        qty = _qty(_cell(row, columns, "qty"))
        if day is None or qty <= 0:
            invalid += 1
            _record_error(
                errors,
                offset,
                "needs a valid date and a quantity above "
                "zero.",
            )
            continue

        item_no = _text(_cell(row, columns, "no"))
        description = _text(_cell(row, columns, "desc"))
        item = by_no.get(item_no) if item_no else None
        if item is None and description:
            item = by_desc.get(description.lower())
        if item is None:
            item_not_found += 1
            _record_error(
                errors, offset, "item not found in this site's list."
            )
            continue

        if item.is_heading:
            invalid += 1
            _record_error(
                errors,
                offset,
                "that row is a group - use the item under it.",
            )
            continue

        if day not in locked_cache:
            locked_cache[day] = not dpr_service.is_day_editable(
                site, day, today=today
            )
        if locked_cache[day]:
            skipped_locked += 1
            continue

        location = _text(_cell(row, columns, "location"))
        agency = _text(_cell(row, columns, "agency"))
        remarks = _text(_cell(row, columns, "remark"))
        key = (item.id, day, qty, location, agency, remarks)
        if existing_counter[key] > 0:
            existing_counter[key] -= 1
            skipped_duplicate += 1
            continue

        DprEntry.objects.create(
            item=item,
            date=day,
            qty=qty,
            rate=boq.effective_rate(item, day, series),
            location=location,
            agency=agency,
            remarks=remarks,
            source=DprEntrySource.EXCEL,
            created_by=actor,
            updated_by=actor,
        )
        created += 1

    return {
        "created": created,
        "skipped_duplicate": skipped_duplicate,
        "skipped_locked": skipped_locked,
        "item_not_found": item_not_found,
        "invalid": invalid,
        "errors": errors,
    }


def build_item_template() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "DPR items"
    sheet.append(
        [
            "Item no",
            "Description",
            "Unit",
            "Qty",
            "Authority rate",
            "Tender % (+ above / - below)",
            "Quoted rate",
        ]
    )
    # A group has a number and a name only; the rows nested under it
    # (4.1, 4.2 ...) are its items. Leave Tender % blank to use the
    # contract-wide percentage; give Quoted rate instead of a
    # percentage if that is what the tender sheet shows.
    sheet.append(["4", "Earthwork"])
    sheet.append(
        ["4.1", "Earthwork in embankment", "cum", 1000, 250, -5.5]
    )
    sheet.append(
        ["4.2", "Blanketing", "cum", 400, 900, None, 936]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_dpr_template(site, today=None) -> bytes:
    today = today or timezone.localdate()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "DPR"
    sheet.append(
        [
            "Date (YYYY-MM-DD)",
            "Item no",
            "Description",
            "Unit",
            "Qty done",
            "Location / chainage",
            "Agency",
            "Remarks",
        ]
    )
    for item in DprItem.objects.filter(
        site=site, is_active=True
    ):
        sheet.append(
            [
                today.isoformat(),
                item.item_no,
                item.description,
                item.unit,
                None,
                "",
                "",
                "",
            ]
        )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
