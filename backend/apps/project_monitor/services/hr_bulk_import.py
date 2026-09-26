"""
Company-wide HR uploads: the staff register and the muster.

HR keeps these in Excel for every site at once, so one file carries a
**Site code** column and each row lands on its own site (the same
routing ``hr_import`` uses for labour rows). Every row is checked
against the uploader's right to enter HR data for that row's site -
for the HR Department that is every site.

Staff register - ``Site code | Staff code | Name | Designation |
Monthly salary | From date | To date | Effective from``:

- A person not yet on the site's register is added.
- The same person (matched by site + staff code, else site + name)
  with the same salary changes nothing; a new Designation, To date
  (they left) or a first Staff code is filled in.
- A **different salary** is never applied silently - editing it
  in place would rewrite past days. It is reported as a conflict
  unless the row carries an ``Effective from`` date, in which case the
  old row ends the day before and a new row starts at that salary.

Muster (attendance) - two layouts, detected automatically:

- flat rows: ``Site code | Staff code | Staff name | Date | Status``;
- month grid: one row per person, one column per day (real dates, or
  day numbers 1-31 with the month named in the sheet or given).

A day is Present (normal cost), Half day (half the day's rate) or
Absent (0). Only Absent and Half day create a day override; Present
clears an override the muster made earlier but never a cost someone
entered by hand. Re-uploading the same file changes nothing.
"""

import calendar
import io
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal

import openpyxl
from django.core.exceptions import (
    ValidationError as DjangoValidationError,
)
from django.db import transaction
from django.utils import timezone
from rest_framework.exceptions import ValidationError

from apps.organization.models import Site
from apps.project_monitor.models import (
    StaffDayOverride,
    StaffMember,
)
from apps.project_monitor.services import hr, site_access
from apps.project_monitor.services.dpr_import import (
    HEADER_SCAN_ROWS,
    _cell,
    _detect_columns,
    _money,
    _text,
    _to_date,
    read_rows,
)

MAX_REPORTED_ERRORS = 100
MUSTER_NOTE = "Muster"

_CODE = (
    r"code|employee\s*id|staff\s*id|emp\s*(id|no|code)|^id$"
)

STAFF_PATTERNS = [
    ("project", r"project|site"),
    ("code", _CODE),
    ("effective", r"effective|revis|increment"),
    ("from", r"from|join|start|doj|w\.?e\.?f"),
    ("to", r"^to\b|to date|end|leav|relie|left|dol|last"),
    ("designation", r"design|role|post|position|title"),
    ("salary", r"salary|monthly|ctc|pay|gross"),
    ("name", r"name"),
]
STAFF_TEMPLATE_HEADER = [
    "Site code",
    "Staff code",
    "Name",
    "Designation",
    "Monthly salary",
    "From date",
    "To date",
    "Effective from",
]

FLAT_PATTERNS = [
    ("project", r"project|site"),
    ("code", _CODE),
    ("date", r"date|day"),
    ("status", r"status|attend|present|mark|p/a"),
    ("name", r"name|staff|employee"),
]
GRID_IDENTITY = [
    ("project", r"project|site"),
    ("code", _CODE),
    ("name", r"name|staff|employee"),
]
MUSTER_FLAT_HEADER = [
    "Site code",
    "Staff code",
    "Staff name",
    "Date",
    "Status",
]

PRESENT = "P"
HALF = "H"
ABSENT = "A"
_STATUS_WORDS = {
    PRESENT: {
        "p", "pr", "present", "w", "wo", "w/o", "off", "weekly off",
        "h", "hol", "holiday", "l", "pl", "cl", "el", "sl", "leave",
        "od", "on duty", "1", "1.0", "full", "full day", "fd",
    },
    HALF: {
        "hd", "h/d", "half", "half day", "halfday", "0.5", "½",
        "p/2", "1/2",
    },
    ABSENT: {
        "a", "ab", "absent", "lwp", "ul", "unpaid", "lop", "0",
        "0.0", "nil",
    },
}
STATUS_LEGEND = [
    ("P", "Present - a normal day (also W/O weekly off, H holiday, "
     "L/PL/CL/SL/EL paid leave, OD on duty)"),
    ("HD", "Half day - half of the day's rate"),
    ("A", "Absent - costs nothing (also AB, LWP, LOP, UL unpaid)"),
    ("(blank)", "No entry - leaves that day as it is"),
]

_MONTH_NUMBERS = {
    name.lower(): number
    for number, name in enumerate(calendar.month_name)
    if name
}
_MONTH_NUMBERS.update(
    {
        name.lower(): number
        for number, name in enumerate(calendar.month_abbr)
        if name
    }
)
_MONTH_NUMBERS["sept"] = 9


# ------------------------------------------------------------ shared


class _Result:
    """Counters, a per-site breakdown and a capped error log."""

    def __init__(self, counters):
        self.counters = dict.fromkeys(counters, 0)
        self.by_site = defaultdict(lambda: defaultdict(int))
        self.errors = []
        self.error_count = 0

    def count(self, key, site=None):
        self.counters[key] += 1
        if site is not None:
            self.by_site[site.site_code][key] += 1

    def note(self, number, message):
        """An entry in the error log that is not its own counter."""
        self.error_count += 1
        if len(self.errors) < MAX_REPORTED_ERRORS:
            self.errors.append(f"Row {number}: {message}")

    def reject(self, key, number, message, site=None):
        self.count(key, site)
        self.note(number, message)

    def as_dict(self) -> dict:
        return {
            **self.counters,
            "errors": self.errors,
            "error_count": self.error_count,
            "by_site": [
                {"site_code": code, **dict(counts)}
                for code, counts in sorted(self.by_site.items())
            ],
        }


class _SiteRouter:
    """Row -> site, with the uploader's HR right checked per site."""

    def __init__(self, user, default_site):
        self.user = user
        self.default_site = default_site
        self._sites: dict = {}
        self._allowed: dict = {}

    def route(self, code):
        """``(site, None)`` or ``(None, (counter, message))``."""
        if code:
            key = code.lower()
            if key not in self._sites:
                self._sites[key] = Site.objects.filter(
                    site_code__iexact=code
                ).first()
            site = self._sites[key]
            if site is None:
                return None, (
                    "site_not_found",
                    f'no site has the code "{code}".',
                )
        elif self.default_site is not None:
            site = self.default_site
        else:
            return None, ("invalid", "the Site code is blank.")

        if site.id not in self._allowed:
            self._allowed[site.id] = site_access.can_enter_hr(
                self.user, site
            )
        if not self._allowed[site.id]:
            return None, (
                "not_permitted",
                f"you cannot enter HR data for {site.site_code}.",
            )
        return site, None


class _StaffIndex:
    """A site's register, read once and kept current as rows add."""

    def __init__(self):
        self._by_site: dict = {}

    def rows(self, site):
        if site.id not in self._by_site:
            self._by_site[site.id] = list(
                StaffMember.objects.filter(site=site)
            )
        return self._by_site[site.id]

    def add(self, site, member):
        self.rows(site).append(member)

    def matches(self, site, code, name):
        """Every register row for this person (a person can have
        several after salary revisions)."""
        rows = self.rows(site)
        if code:
            found = [
                m for m in rows if m.staff_code.lower() == code.lower()
            ]
            if found:
                return found
            # The register may not carry codes yet: fall back to a
            # code-less row with the same name.
            return [
                m
                for m in rows
                if not m.staff_code
                and name
                and m.name.lower() == name.lower()
            ]
        return [
            m
            for m in rows
            if name and m.name.lower() == name.lower()
        ]

    def latest(self, site, code, name):
        found = self.matches(site, code, name)
        return max(found, key=lambda m: m.from_date) if found else None


def _blank_row(row) -> bool:
    return not any(_text(cell) for cell in row)


def _optional_date(cell):
    """``(date_or_None, ok)``: blank is fine, junk is not."""
    if not _text(cell):
        return None, True
    parsed = _to_date(cell)
    return parsed, parsed is not None


def _validation_message(error: DjangoValidationError) -> str:
    if hasattr(error, "message_dict"):
        parts = [
            " ".join(messages)
            for messages in error.message_dict.values()
        ]
        return " ".join(parts)
    return " ".join(error.messages)


def _save(member) -> None:
    member.full_clean()
    member.save()


# ---------------------------------------------------- staff register


STAFF_COUNTERS = (
    "staff_created",
    "staff_updated",
    "staff_revised",
    "staff_unchanged",
    "salary_conflict",
    "site_not_found",
    "not_permitted",
    "invalid",
)


def _find_staff_header(rows, has_default_site):
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        columns = _detect_columns(row, STAFF_PATTERNS)
        if (
            "name" in columns
            and "salary" in columns
            and "from" in columns
            and ("project" in columns or has_default_site)
        ):
            return index, columns
    return None, None


@transaction.atomic
def import_staff_register(
    uploaded_file, user, default_site=None, today=None
) -> dict:
    today = today or timezone.localdate()
    rows = read_rows(uploaded_file)
    header_index, columns = _find_staff_header(
        rows, default_site is not None
    )
    if header_index is None:
        raise ValidationError(
            {
                "file": (
                    "Could not find a header row with Site code, "
                    "Name, Monthly salary and From date columns."
                )
            }
        )

    router = _SiteRouter(user, default_site)
    index = _StaffIndex()
    result = _Result(STAFF_COUNTERS)

    for number, row in enumerate(
        rows[header_index + 1 :], start=header_index + 2
    ):
        if _blank_row(row):
            continue

        site, problem = router.route(
            _text(_cell(row, columns, "project"))
        )
        if problem:
            result.reject(problem[0], number, problem[1])
            continue

        name = _text(_cell(row, columns, "name"))
        code = _text(_cell(row, columns, "code"))
        designation = _text(_cell(row, columns, "designation"))
        if not name:
            result.reject(
                "invalid", number, "the Name is blank.", site
            )
            continue

        salary_cell = _cell(row, columns, "salary")
        if not _text(salary_cell):
            result.reject(
                "invalid",
                number,
                f"{name}: the Monthly salary is blank.",
                site,
            )
            continue
        salary = _money(salary_cell)
        if salary < 0:
            result.reject(
                "invalid",
                number,
                f"{name}: the salary cannot be negative.",
                site,
            )
            continue

        from_date = _to_date(_cell(row, columns, "from"))
        if from_date is None:
            result.reject(
                "invalid",
                number,
                f"{name}: needs a valid From date.",
                site,
            )
            continue
        to_date, to_ok = _optional_date(
            _cell(row, columns, "to")
        )
        effective, effective_ok = _optional_date(
            _cell(row, columns, "effective")
        )
        if not to_ok or not effective_ok:
            result.reject(
                "invalid",
                number,
                f"{name}: a To date / Effective from date is "
                "not a valid date.",
                site,
            )
            continue

        person = index.latest(site, code, name)
        try:
            # A savepoint per row: a revision that fails halfway
            # leaves nothing behind.
            with transaction.atomic():
                outcome = _apply_staff_row(
                    site=site,
                    index=index,
                    person=person,
                    code=code,
                    name=name,
                    designation=designation,
                    salary=salary,
                    from_date=from_date,
                    to_date=to_date,
                    effective=effective,
                    user=user,
                )
        except DjangoValidationError as error:
            result.reject(
                "invalid",
                number,
                f"{name}: {_validation_message(error)}",
                site,
            )
            continue
        if isinstance(outcome, tuple):
            result.reject(outcome[0], number, outcome[1], site)
        else:
            result.count(outcome, site)

    return result.as_dict()


def _apply_staff_row(
    *,
    site,
    index,
    person,
    code,
    name,
    designation,
    salary,
    from_date,
    to_date,
    effective,
    user,
):
    """A counter name on success, or ``(counter, message)`` to reject."""
    if person is None:
        member = StaffMember(
            site=site,
            staff_code=code,
            name=name,
            designation=designation,
            monthly_salary=salary,
            from_date=from_date,
            to_date=to_date,
            created_by=user,
            updated_by=user,
        )
        _save(member)
        index.add(site, member)
        return "staff_created"

    if salary != person.monthly_salary:
        if effective is None:
            return (
                "salary_conflict",
                f"{name}: the salary in the file ({salary}) is not "
                f"the register's ({person.monthly_salary}). Add an "
                '"Effective from" date to revise it.',
            )
        if effective <= person.from_date:
            return (
                "invalid",
                f"{name}: Effective from must be after their "
                f"start date ({person.from_date}).",
            )
        if person.to_date is not None and effective > person.to_date:
            return (
                "invalid",
                f"{name}: they left on {person.to_date}, before "
                "the Effective from date.",
            )
        new_to = to_date if to_date is not None else person.to_date
        person.to_date = effective - timedelta(days=1)
        person.updated_by = user
        _save(person)
        revised = StaffMember(
            site=site,
            staff_code=code or person.staff_code,
            name=person.name,
            designation=designation or person.designation,
            monthly_salary=salary,
            from_date=effective,
            to_date=new_to,
            created_by=user,
            updated_by=user,
        )
        _save(revised)
        index.add(site, revised)
        return "staff_revised"

    changed = False
    if designation and designation != person.designation:
        person.designation = designation
        changed = True
    if to_date is not None and to_date != person.to_date:
        person.to_date = to_date
        changed = True
    if code and not person.staff_code:
        person.staff_code = code
        changed = True
    if not changed:
        return "staff_unchanged"
    person.updated_by = user
    _save(person)
    return "staff_updated"


# ------------------------------------------------------------ muster


MUSTER_COUNTERS = (
    "muster_absent",
    "muster_half",
    "muster_cleared",
    "muster_unchanged",
    "muster_kept_manual",
    "staff_not_found",
    "off_payroll",
    "site_not_found",
    "not_permitted",
    "invalid",
)


def parse_status(value):
    """``"P"``/``"H"``/``"A"``, ``None`` for blank, ``"?"`` unknown."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "?"
    if isinstance(value, (int, float, Decimal)):
        text = str(value).lower()
    else:
        text = _text(value).lower()
    if not text:
        return None
    for status, words in _STATUS_WORDS.items():
        if text in words:
            return status
    return "?"


def _month_in_text(value):
    """
    ``(year, month)`` from a cell such as ``September 2026``,
    ``Sep-26``, ``2026-09`` or a title like ``Muster for Sep 2026``.
    """
    if isinstance(value, (date, datetime)):
        return value.year, value.month
    text = _text(value)
    if not text:
        return None

    named = re.search(
        r"\b([A-Za-z]{3,9})\.?[\s,\-/]+(\d{4}|\d{2})\b", text
    )
    if named:
        month = _MONTH_NUMBERS.get(named.group(1).lower())
        if month:
            year = int(named.group(2))
            return (year if year > 99 else 2000 + year), month
    iso = re.search(r"\b(\d{4})-(\d{1,2})\b", text)
    if iso and 1 <= int(iso.group(2)) <= 12:
        return int(iso.group(1)), int(iso.group(2))
    slash = re.search(r"\b(\d{1,2})[/-](\d{4})\b", text)
    if slash and 1 <= int(slash.group(1)) <= 12:
        return int(slash.group(2)), int(slash.group(1))
    return None


def _grid_date_columns(row, taken):
    """``[(column, "date"|"day", value)]`` for the day columns."""
    found = []
    for column, cell in enumerate(row):
        if column in taken or cell is None:
            continue
        if isinstance(cell, (date, datetime)):
            day = cell.date() if isinstance(cell, datetime) else cell
            found.append((column, "date", day))
            continue
        text = _text(cell)
        if not text:
            continue
        parsed = _to_date(text)
        if parsed is not None:
            found.append((column, "date", parsed))
            continue
        if re.fullmatch(r"\d{1,2}", text) and 1 <= int(text) <= 31:
            found.append((column, "day", int(text)))
    return found


def _find_muster_layout(rows, has_default_site):
    """``(kind, header_index, columns, date_columns)`` or ``None``."""
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        flat = _detect_columns(row, FLAT_PATTERNS)
        if (
            "date" in flat
            and "status" in flat
            and ("name" in flat or "code" in flat)
            and ("project" in flat or has_default_site)
        ):
            return "flat", index, flat, None
        identity = _detect_columns(row, GRID_IDENTITY)
        if ("name" in identity or "code" in identity) and (
            "project" in identity or has_default_site
        ):
            columns = _grid_date_columns(
                row, set(identity.values())
            )
            if len(columns) >= 2:
                return "grid", index, identity, columns
    return None


class _OverrideCache:
    """A site's day overrides for a month, read once."""

    def __init__(self):
        self._months: dict = {}

    def _load(self, site, year, month):
        key = (site.id, year, month)
        if key not in self._months:
            start, end = hr.month_bounds(year, month)
            self._months[key] = {
                (row.staff_id, row.date): row
                for row in StaffDayOverride.objects.filter(
                    staff__site=site, date__range=(start, end)
                )
            }
        return self._months[key]

    def get(self, site, staff_id, day):
        return self._load(site, day.year, day.month).get(
            (staff_id, day)
        )

    def put(self, site, row):
        self._load(site, row.date.year, row.date.month)[
            (row.staff_id, row.date)
        ] = row

    def forget(self, site, row):
        self._load(site, row.date.year, row.date.month).pop(
            (row.staff_id, row.date), None
        )


def _apply_status(
    *, site, member, day, status, cache, result, number, user, today
):
    row = cache.get(site, member.id, day)
    if status == PRESENT:
        if row is None:
            result.count("muster_unchanged", site)
        elif (row.note or "").startswith(MUSTER_NOTE):
            row.delete()
            cache.forget(site, row)
            result.count("muster_cleared", site)
        else:
            # A cost someone entered by hand is not the muster's to
            # remove.
            result.count("muster_kept_manual", site)
        return

    if status == ABSENT:
        target = Decimal("0.00")
        note, counter = f"{MUSTER_NOTE}: absent", "muster_absent"
    else:
        target = _money(hr.daily_rate(member, day) / 2)
        note, counter = f"{MUSTER_NOTE}: half day", "muster_half"
    if row is not None and row.amount == target:
        result.count("muster_unchanged", site)
        return
    saved, _ = hr.set_override(
        staff=member,
        day=day,
        amount=target,
        note=note,
        actor=user,
        today=today,
    )
    cache.put(site, saved)
    result.count(counter, site)


def _member_for_day(index, site, code, name, day):
    """``(member, problem)`` for a person on a given day."""
    found = index.matches(site, code, name)
    if not found:
        return None, "not_found"
    covering = [m for m in found if hr.is_on_payroll(m, day)]
    if not covering:
        return None, "off_payroll"
    if len(covering) > 1:
        return None, "ambiguous"
    return covering[0], None


@transaction.atomic
def import_muster(
    uploaded_file,
    user,
    default_site=None,
    month=None,
    today=None,
) -> dict:
    today = today or timezone.localdate()
    rows = read_rows(uploaded_file)
    layout = _find_muster_layout(rows, default_site is not None)
    if layout is None:
        raise ValidationError(
            {
                "file": (
                    "Could not recognise the muster. Use either "
                    "columns Site code, Staff name, Date and Status "
                    "(one row per person per day), or one row per "
                    "person with a column for each day."
                )
            }
        )
    kind, header_index, columns, date_columns = layout

    router = _SiteRouter(user, default_site)
    index = _StaffIndex()
    cache = _OverrideCache()
    result = _Result(MUSTER_COUNTERS)

    def find_and_apply(
        site, code, name, day, status, number, reported
    ):
        member, problem = _member_for_day(
            index, site, code, name, day
        )
        who = name or code
        if problem == "not_found":
            if not reported.get("not_found"):
                reported["not_found"] = True
                result.reject(
                    "staff_not_found",
                    number,
                    f'"{who}" is not on {site.site_code}\'s staff '
                    "register.",
                    site,
                )
            return
        if problem == "off_payroll":
            result.count("off_payroll", site)
            if not reported.get("off_payroll"):
                reported["off_payroll"] = True
                result.note(
                    number,
                    f'"{who}" was not on the payroll on {day} '
                    "(and possibly other days in this row).",
                )
            return
        if problem == "ambiguous":
            result.reject(
                "invalid",
                number,
                f'more than one "{who}" on {site.site_code}\'s '
                "register on that day - use the staff code.",
                site,
            )
            return
        try:
            _apply_status(
                site=site,
                member=member,
                day=day,
                status=status,
                cache=cache,
                result=result,
                number=number,
                user=user,
                today=today,
            )
        except ValidationError as error:
            result.reject(
                "invalid",
                number,
                f"{who} on {day}: {error.detail}",
                site,
            )

    if kind == "flat":
        _process_flat(
            rows,
            header_index,
            columns,
            router,
            result,
            find_and_apply,
            today,
        )
    else:
        year_month = _resolve_month(
            rows, header_index, columns, date_columns, month, today
        )
        _process_grid(
            rows,
            header_index,
            columns,
            date_columns,
            year_month,
            router,
            result,
            find_and_apply,
            today,
        )
    return result.as_dict()


def _process_flat(
    rows, header_index, columns, router, result, apply, today
):
    for number, row in enumerate(
        rows[header_index + 1 :], start=header_index + 2
    ):
        if _blank_row(row):
            continue
        site, problem = router.route(
            _text(_cell(row, columns, "project"))
        )
        if problem:
            result.reject(problem[0], number, problem[1])
            continue

        name = _text(_cell(row, columns, "name"))
        code = _text(_cell(row, columns, "code"))
        if not name and not code:
            result.reject(
                "invalid",
                number,
                "the Staff name / code is blank.",
                site,
            )
            continue
        day = _to_date(_cell(row, columns, "date"))
        if day is None or day > today:
            result.reject(
                "invalid",
                number,
                "needs a valid date that is not in the future.",
                site,
            )
            continue
        status = parse_status(_cell(row, columns, "status"))
        if status is None or status == "?":
            result.reject(
                "invalid",
                number,
                f'the Status "{_text(_cell(row, columns, "status"))}" '
                "is not P, HD or A.",
                site,
            )
            continue
        apply(site, code, name, day, status, number, {})


def _resolve_month(
    rows, header_index, columns, date_columns, month, today
):
    """The (year, month) day-number columns refer to, if any do."""
    if not any(kind == "day" for _c, kind, _v in date_columns):
        return None
    if month:
        year, mon = hr.parse_month(month)
        return year, mon
    taken = {column for column, _k, _v in date_columns}
    for row in rows[: header_index + 1]:
        for column, cell in enumerate(row):
            if column in taken:
                continue
            found = _month_in_text(cell)
            if found:
                return found
    raise ValidationError(
        {
            "file": (
                "The sheet numbers its days 1-31 but does not say "
                "which month. Pick the month, or put it (for "
                "example September 2026) above the header."
            )
        }
    )


def _process_grid(
    rows,
    header_index,
    columns,
    date_columns,
    year_month,
    router,
    result,
    apply,
    today,
):
    resolved = []
    for column, kind, value in date_columns:
        if kind == "date":
            resolved.append((column, value))
            continue
        year, mon = year_month
        try:
            resolved.append((column, date(year, mon, value)))
        except ValueError:
            continue  # e.g. 31 in a 30-day month

    for number, row in enumerate(
        rows[header_index + 1 :], start=header_index + 2
    ):
        if _blank_row(row):
            continue
        name = _text(_cell(row, columns, "name"))
        code = _text(_cell(row, columns, "code"))
        if not name and not code:
            continue  # totals / separator rows
        site, problem = router.route(
            _text(_cell(row, columns, "project"))
        )
        if problem:
            result.reject(problem[0], number, problem[1])
            continue

        reported: dict = {}
        for column, day in resolved:
            cell = row[column] if column < len(row) else None
            status = parse_status(cell)
            if status is None:
                continue
            if status == "?":
                result.reject(
                    "invalid",
                    number,
                    f'{name or code}: "{_text(cell)}" on {day} is '
                    "not P, HD or A.",
                    site,
                )
                continue
            if day > today:
                result.reject(
                    "invalid",
                    number,
                    f"{name or code}: {day} is in the future.",
                    site,
                )
                continue
            apply(site, code, name, day, status, number, reported)


# ---------------------------------------------------------- templates


def _site_codes_sheet(workbook):
    sheet = workbook.create_sheet("Site codes")
    sheet.append(["Site code", "Site name"])
    for site in Site.objects.filter(is_active=True).order_by(
        "site_code"
    ):
        sheet.append([site.site_code, site.site_name])


def _workbook_bytes(workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_staff_template(site=None) -> bytes:
    code = site.site_code if site else "SITE01"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Staff register"
    sheet.append(STAFF_TEMPLATE_HEADER)
    sheet.append(
        [
            code,
            "EMP001",
            "Site Engineer name",
            "Site Engineer",
            42000,
            "2026-04-01",
            None,
            None,
        ]
    )
    sheet.append(
        [
            code,
            "EMP002",
            "Store Keeper name",
            "Store Keeper",
            24000,
            "2026-04-01",
            "2026-09-30",
            None,
        ]
    )
    sheet.append(
        [
            code,
            "EMP001",
            "Site Engineer name",
            "Senior Site Engineer",
            48000,
            "2026-04-01",
            None,
            "2026-10-01",
        ]
    )
    notes = workbook.create_sheet("How to fill")
    for line in (
        "One row per person per site. Site code must match a code "
        "on the Site codes sheet.",
        "A person already on the register with the same salary is "
        "left as it is (a new Designation, To date or Staff code "
        "is filled in).",
        "To change a salary, repeat the person with the new salary "
        "and an Effective from date - the old salary ends the day "
        "before. Without that date a different salary is refused.",
        "Dates: YYYY-MM-DD or DD-MM-YYYY.",
    ):
        notes.append([line])
    _site_codes_sheet(workbook)
    return _workbook_bytes(workbook)


def build_muster_template(
    kind="flat", month=None, site=None, today=None
) -> bytes:
    today = today or timezone.localdate()
    year, mon = hr.parse_month(month, today)
    code = site.site_code if site else "SITE01"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    if kind == "grid":
        sheet.title = "Muster"
        days = hr.month_bounds(year, mon)[1].day
        sheet.append([f"Muster for {date(year, mon, 1):%B %Y}"])
        sheet.append(
            ["Site code", "Staff code", "Staff name"]
            + list(range(1, days + 1))
        )
        for staff_code, name, marks in (
            ("EMP001", "Site Engineer name", {3: "A", 10: "HD"}),
            ("EMP002", "Store Keeper name", {}),
        ):
            sheet.append(
                [code, staff_code, name]
                + [
                    marks.get(day, "P")
                    for day in range(1, days + 1)
                ]
            )
    else:
        sheet.title = "Muster"
        sheet.append(MUSTER_FLAT_HEADER)
        day = date(year, mon, min(3, hr.month_bounds(year, mon)[1].day))
        sheet.append([code, "EMP001", "Site Engineer name", day.isoformat(), "A"])
        sheet.append([code, "EMP001", "Site Engineer name", (day + timedelta(days=7)).isoformat(), "HD"])
    legend = workbook.create_sheet("Status codes")
    legend.append(["Code", "Meaning"])
    for status_code, meaning in STATUS_LEGEND:
        legend.append([status_code, meaning])
    _site_codes_sheet(workbook)
    return _workbook_bytes(workbook)
